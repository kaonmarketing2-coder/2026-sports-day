import sqlite3
from config import DB_PATH

_NONE_STRINGS = {"none", "null", "nan", "", "-"}

def _clean(v):
    """None·'None'·'none'·빈값 → None, 나머지는 strip된 문자열"""
    if v is None:
        return None
    s = str(v).strip()
    return None if s.lower() in _NONE_STRINGS else s

def _clean_emp(d):
    """직원 dict의 선택 컬럼 정규화"""
    opt = ("본부", "그룹", "부서팀", "파트")
    return {k: (_clean(v) if k in opt else v) for k, v in d.items()}


def get_conn():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = get_conn()
    c = conn.cursor()
    c.executescript("""
        CREATE TABLE IF NOT EXISTS employees (
            사번 TEXT PRIMARY KEY,
            이름 TEXT NOT NULL,
            소속 TEXT,
            직책 TEXT,
            본부 TEXT,
            그룹 TEXT,
            부서팀 TEXT,
            파트 TEXT,
            입사일 TEXT,
            상태 TEXT DEFAULT '재직',
            퇴사일 TEXT
        );

        CREATE TABLE IF NOT EXISTS wellmate (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            멘토_이름 TEXT NOT NULL,
            멘티_사번 TEXT,
            마감월 TEXT,
            엔드서베이 TEXT,
            재직확인 TEXT,
            퇴사일 TEXT,
            메모 TEXT,
            생성일 TEXT DEFAULT (date('now'))
        );

        CREATE TABLE IF NOT EXISTS budget_checks (
            ym TEXT PRIMARY KEY
        );
    """)
    # 기존 DB에 컬럼 추가 (마이그레이션)
    for col, default in [("상태", "'재직'"), ("퇴사일", "NULL"), ("메모", "NULL")]:
        try:
            c.execute(f"ALTER TABLE employees ADD COLUMN {col} TEXT DEFAULT {default}")
        except Exception:
            pass
    conn.commit()
    conn.close()


def import_employees_from_df(df):
    conn = get_conn()
    c = conn.cursor()
    count = 0
    for _, row in df.iterrows():
        sabun = str(row.get("사번", "")).strip()
        name = str(row.get("이름", "")).strip()
        if not sabun or not name or sabun == "nan":
            continue
        c.execute("""
            INSERT INTO employees (사번, 이름, 소속, 직책, 본부, 그룹, 부서팀, 파트, 입사일, 상태)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, '재직')
            ON CONFLICT(사번) DO UPDATE SET
                이름=excluded.이름, 소속=excluded.소속, 직책=excluded.직책,
                본부=excluded.본부, 그룹=excluded.그룹, 부서팀=excluded.부서팀,
                파트=excluded.파트, 입사일=excluded.입사일
        """, (
            sabun, name,
            _val(row, "소속"), _val(row, "직책"), _val(row, "본부"),
            _val(row, "그룹"), _val(row, "부서팀"), _val(row, "파트"),
            _val(row, "입사일"),
        ))
        count += 1
    conn.commit()
    conn.close()
    return count


def _val(row, col):
    v = row.get(col, None)
    if v is None:
        return None
    import pandas as pd
    try:
        if pd.isna(v):
            return None
    except Exception:
        pass
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    return None if s in ("nan", "None", "-", "") else s


def get_employee(sabun):
    conn = get_conn()
    row = conn.execute("SELECT * FROM employees WHERE 사번=?", (sabun,)).fetchone()
    conn.close()
    return _clean_emp(dict(row)) if row else None


def search_employees(q, include_resigned=False):
    conn = get_conn()
    status_filter = "" if include_resigned else "AND (상태 IS NULL OR 상태='재직')"
    rows = conn.execute(
        f"SELECT * FROM employees WHERE (이름 LIKE ? OR 사번 LIKE ?) {status_filter} LIMIT 20",
        (f"%{q}%", f"%{q}%")
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def get_all_employees(include_resigned=False):
    conn = get_conn()
    if include_resigned:
        rows = conn.execute("SELECT * FROM employees ORDER BY 입사일 DESC").fetchall()
    else:
        rows = conn.execute("SELECT * FROM employees WHERE (상태 IS NULL OR 상태='재직') ORDER BY 입사일 DESC").fetchall()
    conn.close()
    return [_clean_emp(dict(r)) for r in rows]


def get_resigned_employees():
    conn = get_conn()
    rows = conn.execute("SELECT * FROM employees WHERE 상태='퇴사' ORDER BY 퇴사일 DESC").fetchall()
    conn.close()
    return [dict(r) for r in rows]


def add_employee(data):
    conn = get_conn()
    conn.execute("""
        INSERT INTO employees (사번, 이름, 소속, 직책, 본부, 그룹, 부서팀, 파트, 입사일, 상태, 메모)
        VALUES (:사번,:이름,:소속,:직책,:본부,:그룹,:부서팀,:파트,:입사일,'재직',:메모)
        ON CONFLICT(사번) DO UPDATE SET
            이름=excluded.이름, 소속=excluded.소속, 직책=excluded.직책,
            본부=excluded.본부, 그룹=excluded.그룹, 부서팀=excluded.부서팀,
            파트=excluded.파트, 입사일=excluded.입사일, 메모=excluded.메모
    """, {**data, "메모": data.get("메모")})
    conn.commit()
    conn.close()


def update_employee(sabun, data):
    conn = get_conn()
    conn.execute("""
        UPDATE employees SET 이름=:이름, 소속=:소속, 직책=:직책, 본부=:본부,
        그룹=:그룹, 부서팀=:부서팀, 파트=:파트, 입사일=:입사일, 메모=:메모 WHERE 사번=:sabun
    """, {**data, "sabun": sabun})
    conn.commit()
    conn.close()


def resign_employee(sabun, resign_date):
    conn = get_conn()
    conn.execute("UPDATE employees SET 상태='퇴사', 퇴사일=? WHERE 사번=?", (resign_date, sabun))
    conn.commit()
    conn.close()


def bulk_update_employees(filter_col, filter_val, update_col, update_val):
    allowed_cols = {"소속", "직책", "본부", "그룹", "부서팀", "파트"}
    if filter_col not in allowed_cols or update_col not in allowed_cols:
        return 0
    conn = get_conn()
    cur = conn.execute(
        f"UPDATE employees SET {update_col}=? WHERE {filter_col}=? AND (상태 IS NULL OR 상태='재직')",
        (update_val, filter_val)
    )
    count = cur.rowcount
    conn.commit()
    conn.close()
    return count


def delete_employee(sabun):
    conn = get_conn()
    conn.execute("DELETE FROM employees WHERE 사번=?", (sabun,))
    conn.commit()
    conn.close()


def get_all_wellmate():
    conn = get_conn()
    rows = conn.execute("SELECT * FROM wellmate ORDER BY id DESC").fetchall()
    conn.close()
    return [dict(r) for r in rows]


def add_wellmate(data):
    conn = get_conn()
    cur = conn.execute("""
        INSERT INTO wellmate (멘토_이름, 멘티_사번, 마감월, 엔드서베이, 재직확인, 퇴사일, 메모)
        VALUES (:멘토_이름,:멘티_사번,:마감월,:엔드서베이,:재직확인,:퇴사일,:메모)
    """, data)
    new_id = cur.lastrowid
    conn.commit()
    conn.close()
    return new_id


def get_pending_wellmates():
    """웰메이트 미배정(멘토_이름이 비어있는) 레코드 + 신규입사자 정보 조인"""
    conn = get_conn()
    rows = conn.execute("""
        SELECT w.id, w.멘티_사번, w.생성일,
               e.이름 as 신규입사자명, e.직책, e.소속, e.부서팀, e.입사일
        FROM wellmate w
        LEFT JOIN employees e ON w.멘티_사번 = e.사번
        WHERE w.멘토_이름 IS NULL OR w.멘토_이름 = ''
        ORDER BY w.id ASC
    """).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def update_wellmate(id, data):
    conn = get_conn()
    conn.execute("""
        UPDATE wellmate SET 멘토_이름=:멘토_이름, 멘티_사번=:멘티_사번,
        마감월=:마감월, 엔드서베이=:엔드서베이, 재직확인=:재직확인,
        퇴사일=:퇴사일, 메모=:메모 WHERE id=:id
    """, {**data, "id": id})
    conn.commit()
    conn.close()


def delete_wellmate(id):
    conn = get_conn()
    conn.execute("DELETE FROM wellmate WHERE id=?", (id,))
    conn.commit()
    conn.close()


def get_wellmate(id):
    conn = get_conn()
    row = conn.execute("SELECT * FROM wellmate WHERE id=?", (id,)).fetchone()
    conn.close()
    return dict(row) if row else None


# 직책 순위 (위에 올수록 상위 직책). 키워드 부분일치로 판정하며 위에서부터 먼저 맞는 것을 채택.
# 예) "책임매니저"는 '매니저'보다 '책임'이 위에 있어 책임으로 분류됨. 순수 "매니저"가 최하위.
# 회사 직책 체계에 맞게 이 목록만 수정하면 정렬 순서가 바뀜.
TITLE_RANK = [
    "대표이사", "대표", "부사장", "부문장", "전무", "상무", "이사대우", "이사",
    "실장", "본부장", "센터장", "그룹장", "담당", "팀장", "파트장",
    "PL", "리더", "수석", "책임", "선임", "전임", "주임", "사원", "매니저",
]


def title_rank(title):
    t = (title or "").strip()
    for i, kw in enumerate(TITLE_RANK):
        if kw and kw in t:
            return i
    # 목록에 없는 직책은 매니저 바로 위(중간)로 둠
    return len(TITLE_RANK) - 1.5


def get_org_tree():
    conn = get_conn()
    rows = conn.execute("""
        SELECT 소속, 본부, 그룹, 부서팀, 이름, 직책
        FROM employees
        WHERE 상태 IS NULL OR 상태='재직'
        ORDER BY 소속, 본부, 그룹, 부서팀, 직책, 이름
    """).fetchall()
    conn.close()

    def empty(v):
        return v is None or str(v).strip() in ("", "-", "nan", "None")

    # 직원을 "자기가 속한 가장 깊은 조직 단계"에 직접 배치.
    # 예) 본부장은 그룹/부서팀이 비어 있으므로 본부 노드에 이름이 달림.
    # unit 구조: {"name", "children": {이름: unit}, "people": [직원...]}
    units = {}  # 최상위(소속) 단위들

    def get_unit(container, name):
        if name not in container:
            container[name] = {"name": name, "children": {}, "people": []}
        return container[name]

    for r in rows:
        levels = []
        for col in ("소속", "본부", "그룹", "부서팀"):
            v = r[col]
            if not empty(v):
                levels.append(str(v).strip())
        if not levels:
            levels = ["미분류"]
        cur = units
        node = None
        for name in levels:
            node = get_unit(cur, name)
            cur = node["children"]
        node["people"].append({"이름": r["이름"], "직책": r["직책"] or ""})

    # 렌더용 재귀 노드 구조로 변환
    # 노드 공통: {"name", "level", "is_person", "children": [...]}
    #  - 조직 노드: "count"(하위 전체 인원수) 포함
    #  - 직원 노드: "title"(직책) 포함, is_person=True
    def conv(unit, level):
        child_nodes = [conv(c, level + 1) for c in unit["children"].values()]
        # 같은 단위의 직원은 상위 직책이 위로 오도록 정렬 (매니저가 최하위)
        sorted_people = sorted(unit["people"], key=lambda p: (title_rank(p["직책"]), p["이름"]))
        person_nodes = [
            {"name": p["이름"], "title": p["직책"], "level": level + 1,
             "is_person": True, "children": []}
            for p in sorted_people
        ]
        count = len(unit["people"]) + sum(c["count"] for c in child_nodes)
        # 해당 단위의 직원(본부장 등)을 먼저, 하위 조직을 뒤에 표시
        children = person_nodes + child_nodes
        return {"name": unit["name"], "count": count, "level": level,
                "is_person": False, "children": children}

    소속_nodes = [conv(u, 1) for u in units.values()]
    if not 소속_nodes:
        return None
    # 최상위는 소속(가온그룹/가온브로드밴드/가온로보틱스 등) 여러 개를 나란히 표시
    return 소속_nodes


# ── 대시보드 통계용 함수들 ────────────────────────────────────────

def _last_n_months(n=12):
    from datetime import date
    result = []
    y, m = date.today().year, date.today().month
    for _ in range(n):
        result.append(f"{y:04d}-{m:02d}")
        m -= 1
        if m == 0:
            m = 12; y -= 1
    return list(reversed(result))


def get_monthly_hire_counts(n=12):
    months = _last_n_months(n)
    conn = get_conn()
    rows = conn.execute("""
        SELECT substr(입사일,1,7) as m, COUNT(*) as c
        FROM employees WHERE 입사일 IS NOT NULL AND 입사일 != ''
        GROUP BY m
    """).fetchall()
    conn.close()
    d = {r["m"]: r["c"] for r in rows}
    return [{"month": m, "count": d.get(m, 0)} for m in months]


def get_monthly_resign_counts(n=12):
    months = _last_n_months(n)
    conn = get_conn()
    rows = conn.execute("""
        SELECT substr(퇴사일,1,7) as m, COUNT(*) as c
        FROM employees WHERE 퇴사일 IS NOT NULL AND 퇴사일 != '' AND 상태='퇴사'
        GROUP BY m
    """).fetchall()
    conn.close()
    d = {r["m"]: r["c"] for r in rows}
    return [{"month": m, "count": d.get(m, 0)} for m in months]


def get_소속_dist():
    conn = get_conn()
    rows = conn.execute("""
        SELECT 소속, COUNT(*) as c FROM employees
        WHERE 상태 IS NULL OR 상태='재직'
        GROUP BY 소속 ORDER BY c DESC
    """).fetchall()
    conn.close()
    return [{"소속": r["소속"] or "미분류", "count": r["c"]} for r in rows]


def get_employees_by_hire_month(yyyymm):
    conn = get_conn()
    rows = conn.execute("""
        SELECT * FROM employees WHERE substr(입사일,1,7)=?
        ORDER BY 입사일, 이름
    """, (yyyymm,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def get_employees_by_resign_month(yyyymm):
    conn = get_conn()
    rows = conn.execute("""
        SELECT * FROM employees WHERE substr(퇴사일,1,7)=? AND 상태='퇴사'
        ORDER BY 퇴사일, 이름
    """, (yyyymm,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def get_employees_by_소속_name(name):
    conn = get_conn()
    rows = conn.execute("""
        SELECT * FROM employees WHERE 소속=?
        AND (상태 IS NULL OR 상태='재직')
        ORDER BY 본부, 그룹, 부서팀, 이름
    """, (name,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]


# 직책 카테고리 분류 (키워드 부분일치, 앞에서부터 먼저 맞는 것 채택)
_TITLE_CATS = [
    ("임원",       ["대표이사","대표","부사장","부문장","전무","상무","이사대우","이사"]),
    ("실장/본부장", ["실장","본부장","센터장"]),
    ("그룹장/담당", ["그룹장","담당"]),
    ("팀장",       ["팀장"]),
    ("파트장/PL",  ["파트장","PL","리더"]),
    ("수석",       ["수석"]),
    ("책임",       ["책임"]),
    ("선임",       ["선임","전임","주임"]),
    ("사원",       ["사원"]),
    ("매니저",     ["매니저"]),
]

def _직책_category(title):
    t = (title or "").strip()
    for cat, kws in _TITLE_CATS:
        if any(kw in t for kw in kws):
            return cat
    return "기타"


def get_직책_dist():
    conn = get_conn()
    rows = conn.execute("""
        SELECT 직책, COUNT(*) as c FROM employees
        WHERE 상태 IS NULL OR 상태='재직'
        GROUP BY 직책
    """).fetchall()
    conn.close()
    groups = {}
    for r in rows:
        cat = _직책_category(r["직책"])
        groups[cat] = groups.get(cat, 0) + r["c"]
    return [{"직책": k, "count": v}
            for k, v in sorted(groups.items(), key=lambda x: -x[1])]


def get_employees_by_직책_cat(cat):
    conn = get_conn()
    rows = conn.execute("""
        SELECT * FROM employees
        WHERE 상태 IS NULL OR 상태='재직'
        ORDER BY 소속, 본부, 이름
    """).fetchall()
    conn.close()
    return [dict(r) for r in rows if _직책_category(r["직책"]) == cat]


def get_avg_tenure():
    from datetime import date, datetime
    conn = get_conn()
    rows = conn.execute("""
        SELECT 입사일 FROM employees
        WHERE (상태 IS NULL OR 상태='재직') AND 입사일 IS NOT NULL AND 입사일 != ''
    """).fetchall()
    conn.close()
    today = date.today()
    days_list = []
    for r in rows:
        try:
            d = datetime.strptime(str(r["입사일"])[:10], "%Y-%m-%d").date()
            days_list.append((today - d).days)
        except Exception:
            pass
    if not days_list:
        return {"years": 0, "months": 0, "text": "-"}
    avg = sum(days_list) // len(days_list)
    years, rem = divmod(avg, 365)
    months = rem // 30
    return {"years": years, "months": months, "text": f"{years}년 {months}개월"}


def get_calendar_events(year, month):
    ym = f"{year:04d}-{month:02d}"
    conn = get_conn()
    rows = conn.execute("""
        SELECT w.id, w.멘토_이름, w.멘티_사번, w.마감월, w.엔드서베이,
               e.이름 as 멘티_이름, e.소속 as 멘티_소속, e.부서팀 as 멘티_팀
        FROM wellmate w
        LEFT JOIN employees e ON w.멘티_사번 = e.사번
        WHERE substr(w.마감월,1,7)=? OR substr(w.엔드서베이,1,7)=?
    """, (ym, ym)).fetchall()
    conn.close()
    events = []
    for r in rows:
        if r["마감월"] and str(r["마감월"])[:7] == ym:
            events.append({
                "date": str(r["마감월"])[:10],
                "type": "마감",
                "멘토": r["멘토_이름"] or "-",
                "멘티": r["멘티_이름"] or r["멘티_사번"] or "-",
                "팀": r["멘티_팀"] or "",
                "소속": r["멘티_소속"] or "-",
                "id": r["id"],
            })
        if r["엔드서베이"] and str(r["엔드서베이"])[:7] == ym:
            events.append({
                "date": str(r["엔드서베이"])[:10],
                "type": "서베이",
                "멘토": r["멘토_이름"] or "-",
                "멘티": r["멘티_이름"] or r["멘티_사번"] or "-",
                "팀": r["멘티_팀"] or "",
                "소속": r["멘티_소속"] or "-",
                "id": r["id"],
            })
    return events


def get_team_list():
    """재직 직원을 팀(부서팀)별로 묶어 반환. 팀 없는 직원은 소속/본부로 묶음."""
    conn = get_conn()
    rows = conn.execute("""
        SELECT 사번, 이름, 직책, 소속, 본부, 그룹, 부서팀
        FROM employees
        WHERE (상태 IS NULL OR 상태='재직')
        ORDER BY 소속, 본부, 그룹, 부서팀, 이름
    """).fetchall()
    conn.close()

    teams = {}
    for r in rows:
        r = _clean_emp(dict(r))
        team_key = r["부서팀"] or f"({r['본부'] or r['소속'] or '미분류'})"
        if team_key not in teams:
            teams[team_key] = {
                "name": r["부서팀"] or team_key,
                "소속": r["소속"] or "-",
                "본부": r["본부"] or "-",
                "그룹": r["그룹"] or "-",
                "members": [],
            }
        teams[team_key]["members"].append({
            "name": r["이름"],
            "title": r["직책"] or "-",
        })

    # 직책 순위로 각 팀 멤버 정렬
    for t in teams.values():
        t["members"].sort(key=lambda m: title_rank(m["title"]))
        t["count"] = len(t["members"])

    return list(teams.values())


def export_to_excel(path: str):
    """직원DB + 웰메이트 현황을 엑셀 파일로 저장한다."""
    import pandas as pd
    from openpyxl.styles import Font, PatternFill
    from datetime import datetime
    conn = get_conn()

    emp_rows = conn.execute("""
        SELECT 사번, 이름, 직책, 소속, 본부, 그룹, 부서팀, 파트, 입사일, 상태, 퇴사일
        FROM employees ORDER BY 상태, 소속, 입사일
    """).fetchall()

    wm_rows = conn.execute("""
        SELECT w.id, w.멘토_이름, w.멘티_사번,
               e.이름 as 멘티_이름, e.소속 as 멘티_소속, e.부서팀 as 멘티_팀, e.직책 as 멘티_직책,
               w.마감월, w.엔드서베이, w.재직확인, w.퇴사일, w.메모, w.생성일
        FROM wellmate w
        LEFT JOIN employees e ON w.멘티_사번 = e.사번
        ORDER BY w.생성일 DESC
    """).fetchall()
    conn.close()

    df_emp = pd.DataFrame([dict(r) for r in emp_rows]) if emp_rows else pd.DataFrame()
    df_wm  = pd.DataFrame([dict(r) for r in wm_rows])  if wm_rows  else pd.DataFrame()

    now_str = datetime.now().strftime("%Y-%m-%d %H:%M")
    gray_fill = PatternFill("solid", fgColor="F3F4F6")
    gray_font = Font(color="6B7280", size=9)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        # 데이터를 2행부터 쓰고, 1행에 수정일시 메모 삽입
        df_emp.to_excel(writer, sheet_name="직원DB",         index=False, startrow=1)
        df_wm.to_excel(writer,  sheet_name="웰메이트배정현황", index=False, startrow=1)

        for sheet_name in ["직원DB", "웰메이트배정현황"]:
            ws = writer.sheets[sheet_name]
            ws.cell(row=1, column=1, value=f"마지막 업데이트: {now_str}")
            cell = ws.cell(row=1, column=1)
            cell.font = gray_font
            cell.fill = gray_fill


def is_budget_checked(ym: str) -> bool:
    conn = get_conn()
    row = conn.execute("SELECT 1 FROM budget_checks WHERE ym=?", (ym,)).fetchone()
    conn.close()
    return row is not None

def set_budget_checked(ym: str):
    conn = get_conn()
    conn.execute("INSERT OR IGNORE INTO budget_checks (ym) VALUES (?)", (ym,))
    conn.commit()
    conn.close()

def unset_budget_checked(ym: str):
    conn = get_conn()
    conn.execute("DELETE FROM budget_checks WHERE ym=?", (ym,))
    conn.commit()
    conn.close()
